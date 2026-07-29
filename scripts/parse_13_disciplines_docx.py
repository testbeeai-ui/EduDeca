"""Parse EduDeca 13-discipline Level 1-3 MCQ docx into seed JSON + SQL."""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document

DOC = Path(r"C:\Users\rentk\Downloads\You listed 13 disciplines Level 1-2 Qs.docx")
OUT_JSON = Path(r"C:\Users\rentk\Desktop\Edublast\EduDeca\_questions_seed.json")
OUT_SQL = Path(r"C:\Users\rentk\Desktop\Edublast\EduDeca\scripts\_seed_questions_l1_l3.sql")

DISC_MAP = {
    "physics": "phy",
    "chemistry": "che",
    "mathematics": "mat",
    "applied mathematics": "amat",
    "biology": "bio",
    "biotechnology": "biotech",
    "computer science and ai": "cs",
    "entrepreneurship": "ent",
    "verbal ability": "eng",
    "quantitative ability": "eco",
    "analytical ability": "log",
    "general knowledge": "gk",
    "financial literacy": "fin",
}

LETTER_TO_INDEX = {"A": 0, "B": 1, "C": 2, "D": 3}


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("\xa0", " ")).strip()


def parse_options(block: str) -> list[str]:
    lines = [clean(x) for x in block.splitlines() if clean(x)]
    opts: list[str] = []
    for line in lines:
        m = re.match(r"^([A-D])\.\s*(.+)$", line)
        if m:
            opts.append(m.group(2).strip())
    if len(opts) != 4:
        raise ValueError(f"Expected 4 options, got {len(opts)}: {block!r}")
    return opts


def parse_questions(doc: Document) -> list[dict]:
    paras = [clean(p.text) for p in doc.paragraphs]
    # Collapse consecutive option lines already in one para; keep empties as separators
    raw = [(p.text or "").rstrip() for p in doc.paragraphs]

    questions: list[dict] = []
    level = None
    i = 0
    while i < len(raw):
        line = clean(raw[i])
        if re.match(r"^SET\s+\d+\s*[—\-–]\s*LEVEL\s+(\d+)", line, re.I):
            level = int(re.search(r"LEVEL\s+(\d+)", line, re.I).group(1))
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.+)$", line)
        if m and level is not None:
            num = int(m.group(1))
            disc_name = clean(m.group(2))
            disc_key = disc_name.lower()
            if disc_key not in DISC_MAP:
                raise ValueError(f"Unknown discipline: {disc_name}")
            subject_id = DISC_MAP[disc_key]

            # stem may be one or more paragraphs until options
            stem_parts: list[str] = []
            i += 1
            while i < len(raw):
                t = raw[i]
                ct = clean(t)
                if re.match(r"^[A-D]\.", ct) or (
                    "\n" in t and re.search(r"^[A-D]\.", t, re.M)
                ):
                    break
                if re.match(r"^(\d+)\.\s+", ct) or re.match(r"^SET\s+\d+", ct, re.I):
                    break
                if re.match(r"^Answer Key", ct, re.I):
                    break
                if ct:
                    stem_parts.append(ct)
                i += 1
            stem = " ".join(stem_parts).strip()
            # Fix duplicated OCR-ish lines in Applied Math L3
            stem = re.sub(
                r"where D is demandP\s*where D is demand and P",
                "where D is demand and P",
                stem,
            )

            if i >= len(raw):
                raise ValueError(f"Missing options for L{level} #{num}")
            opt_block = raw[i]
            # Sometimes options are one para with newlines
            options = parse_options(opt_block)
            i += 1

            questions.append(
                {
                    "level": level,
                    "sort_order": num,
                    "subject_id": subject_id,
                    "discipline": disc_name,
                    "stem": stem,
                    "options": options,
                    "correct_index": None,
                    "explanation": None,
                }
            )
            continue
        i += 1
    return questions


def parse_answer_tables(doc: Document) -> dict[int, list[dict]]:
    """level -> list of {no, letter, explanation?}"""
    out: dict[int, list[dict]] = {}
    for ti, table in enumerate(doc.tables):
        level = ti + 1
        rows = []
        for ri, row in enumerate(table.rows):
            cells = [clean(c.text) for c in row.cells]
            if ri == 0:
                continue
            no = int(cells[0])
            ans = cells[2] if len(cells) > 2 else cells[-1]
            # "B — Ammeter" or "B"
            letter_m = re.match(r"^([A-D])\b", ans)
            if not letter_m:
                raise ValueError(f"Bad answer cell L{level} #{no}: {ans}")
            letter = letter_m.group(1)
            expl = None
            if len(cells) > 3 and cells[3]:
                expl = cells[3]
            rows.append({"no": no, "letter": letter, "explanation": expl})
        out[level] = rows
    return out


def main() -> None:
    doc = Document(str(DOC))
    questions = parse_questions(doc)
    answers = parse_answer_tables(doc)

    by_key: dict[tuple[int, int], dict] = {(q["level"], q["sort_order"]): q for q in questions}

    for level, rows in answers.items():
        for row in rows:
            q = by_key.get((level, row["no"]))
            if not q:
                raise ValueError(f"Answer without question L{level} #{row['no']}")
            q["correct_index"] = LETTER_TO_INDEX[row["letter"]]
            if row["explanation"]:
                q["explanation"] = row["explanation"]

    missing = [q for q in questions if q["correct_index"] is None]
    if missing:
        raise SystemExit(f"Missing answers for {len(missing)} questions")

    # Expect 13 * 3 = 39
    counts: dict[int, int] = {}
    for q in questions:
        counts[q["level"]] = counts.get(q["level"], 0) + 1
    print("counts by level:", counts)
    if counts != {1: 13, 2: 13, 3: 13}:
        raise SystemExit(f"Unexpected counts: {counts}")

    # IDs
    for q in questions:
        q["id"] = f"l{q['level']}-{q['subject_id']}-{q['sort_order']:02d}"
        q["difficulty_rating"] = q["level"]
        q["published"] = True

    OUT_JSON.write_text(json.dumps(questions, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {OUT_JSON} ({len(questions)} questions)")

    # SQL: wipe + insert
    lines = [
        "-- Replace EduDeca question bank with 13-discipline Levels 1-3 pack",
        "BEGIN;",
        "DELETE FROM public.edudeca_questions;",
    ]
    for q in questions:
        opts = json.dumps(q["options"], ensure_ascii=False).replace("'", "''")
        stem = q["stem"].replace("'", "''")
        expl = (q["explanation"] or "").replace("'", "''")
        expl_sql = f"'{expl}'" if q["explanation"] else "NULL"
        lines.append(
            "INSERT INTO public.edudeca_questions "
            "(id, subject_id, level, sort_order, stem, options, correct_index, explanation, difficulty_rating, published) VALUES ("
            f"'{q['id']}', '{q['subject_id']}', {q['level']}, {q['sort_order']}, "
            f"'{stem}', '{opts}'::jsonb, {q['correct_index']}, {expl_sql}, {q['difficulty_rating']}, true);"
        )
    lines.append("COMMIT;")
    OUT_SQL.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Wrote {OUT_SQL}")


if __name__ == "__main__":
    main()
